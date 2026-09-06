#!/usr/bin/env python3
"""Command 1 — build the two demo corpora.

Two very different sources, on purpose, because the point of the knowledge
base is that agentic retrieval routes between them:

  corpus/eu-directives/  EU directives that bear on Spanish employment law,
                         fetched from EUR-Lex by CELEX number. Official texts,
                         published for reuse, and available in several formats
                         from the same URI — which is where most of the format
                         mix comes from.

  corpus/hr-templates/   HR sample documents and templates. Only allowlisted
                         hosts are crawled by default (see ALLOWED_HOSTS):
                         template sites frequently forbid redistribution in
                         their terms, and a demo corpus is still a copy. Add
                         your own with --allow-domain once you've checked the
                         licence — the flag prints a reminder, deliberately.

robots.txt is honoured for every host, requests are rate-limited, and the
User-Agent identifies the demo. Nothing is fabricated: `manifest.json` records
the real URL, HTTP content type, and sniffed format of every file, and the run
summary reports the actual format distribution rather than a target one.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import time
import urllib.robotparser
from dataclasses import asdict, dataclass
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests

from _common import CORPUS_DIR

USER_AGENT = "foundry-iq-demo-scraper/1.0 (+https://learn.microsoft.com/azure/search/agentic-retrieval-overview)"
REQUEST_DELAY_SECONDS = 1.0

# EUR-Lex serves every document under a stable CELEX identifier, so the seed
# list is data rather than a crawl. Ids that 404 are reported and skipped, so a
# stale entry costs a log line instead of a broken run.
EMPLOYMENT_CELEX = [
    "32003L0088",  # Working time
    "32019L1152",  # Transparent and predictable working conditions
    "32019L1158",  # Work-life balance for parents and carers
    "31996L0071",  # Posting of workers
    "32018L0957",  # Posting of workers (amendment)
    "32014L0067",  # Posting of workers (enforcement)
    "31999L0070",  # Fixed-term work
    "31997L0081",  # Part-time work
    "31998L0059",  # Collective redundancies
    "32001L0023",  # Transfer of undertakings
    "32008L0104",  # Temporary agency work
    "32002L0014",  # Information and consultation of employees
    "32009L0038",  # European Works Councils
    "31989L0391",  # Occupational safety and health framework
    "31989L0654",  # Workplace requirements
    "31989L0656",  # Personal protective equipment
    "31990L0269",  # Manual handling of loads
    "31990L0270",  # Display screen equipment
    "32009L0104",  # Work equipment
    "31992L0085",  # Pregnant workers
    "31994L0033",  # Protection of young people at work
    "32000L0043",  # Racial equality
    "32000L0078",  # Equal treatment in employment and occupation
    "32006L0054",  # Equal opportunities and equal treatment (recast)
    "32010L0041",  # Equal treatment, self-employed
    "32023L0970",  # Pay transparency
    "32022L2041",  # Adequate minimum wages
    "31991L0533",  # Written statement of the employment relationship
    "32003L0041",  # Institutions for occupational retirement provision
    "31997L0074",  # European Works Councils (extension to the UK)
]

# Formats EUR-Lex exposes for the same CELEX id. Requesting several is how the
# corpus gets a genuine mix rather than a synthesised one.
EURLEX_FORMATS = [
    ("PDF", "pdf"),
    ("HTML", "html"),
    ("DOC", "doc"),
]

# Hosts whose terms permit reuse of published guidance and templates. Keep this
# conservative; --allow-domain exists for anything you've cleared yourself.
ALLOWED_HOSTS = {
    "www.acas.org.uk",
    "www.gov.uk",
    "assets.publishing.service.gov.uk",
    "europa.eu",
    "eur-lex.europa.eu",
    "osha.europa.eu",
    "www.eurofound.europa.eu",
}

HR_SEED_PAGES = [
    "https://www.acas.org.uk/templates-for-employers",
    "https://www.acas.org.uk/example-written-statement-of-employment-particulars",
    "https://www.acas.org.uk/disciplinary-and-grievance-procedures",
    "https://www.gov.uk/employment-contracts-and-conditions",
    "https://osha.europa.eu/en/publications",
]

DOCUMENT_SUFFIXES = {
    ".pdf": "pdf",
    ".doc": "doc",
    ".docx": "docx",
    ".ppt": "ppt",
    ".pptx": "pptx",
    ".xls": "xls",
    ".xlsx": "xlsx",
    ".txt": "txt",
    ".md": "markdown",
    ".rtf": "rtf",
    ".odt": "odt",
    ".csv": "csv",
}

CONTENT_TYPE_FORMATS = {
    "application/pdf": "pdf",
    "application/msword": "doc",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.ms-powerpoint": "ppt",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.ms-excel": "xls",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/plain": "txt",
    "text/markdown": "markdown",
    "text/csv": "csv",
    "text/html": "html",
}


@dataclass
class Record:
    """One fetched file, as recorded in manifest.json."""

    path: str
    url: str
    corpus: str
    fmt: str
    content_type: str
    bytes: int
    sha256: str
    title: str
    celex: str | None = None
    derived_from: str | None = None


class LinkExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[str] = []
        self.title_parts: list[str] = []
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            href = dict(attrs).get("href")
            if href:
                self.links.append(href)
        elif tag == "title":
            self._in_title = True

    def handle_endtag(self, tag: str) -> None:
        if tag == "title":
            self._in_title = False

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title_parts.append(data)

    @property
    def title(self) -> str:
        return " ".join("".join(self.title_parts).split())


class TextExtractor(HTMLParser):
    """Strips markup to plain text without pulling in a parser dependency."""

    SKIP = {"script", "style", "nav", "footer", "header", "noscript"}

    def __init__(self) -> None:
        super().__init__()
        self.chunks: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self.SKIP:
            self._skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP and self._skip_depth:
            self._skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self._skip_depth and data.strip():
            self.chunks.append(data.strip())

    @property
    def text(self) -> str:
        return re.sub(r"\n{3,}", "\n\n", "\n".join(self.chunks))


class Fetcher:
    """HTTP with robots.txt enforcement and a fixed delay between requests."""

    def __init__(self, allowed_hosts: set[str], *, ignore_robots: bool = False) -> None:
        self.allowed_hosts = allowed_hosts
        self.ignore_robots = ignore_robots
        self.session = requests.Session()
        self.session.headers["User-Agent"] = USER_AGENT
        self._robots: dict[str, urllib.robotparser.RobotFileParser | None] = {}
        self._last_request = 0.0

    def allowed(self, url: str) -> tuple[bool, str]:
        host = urlparse(url).netloc
        if host not in self.allowed_hosts:
            return False, f"host not allowlisted ({host})"
        if self.ignore_robots:
            return True, ""
        parser = self._robots.get(host, "missing")  # type: ignore[assignment]
        if parser == "missing":
            parser = urllib.robotparser.RobotFileParser()
            parser.set_url(f"{urlparse(url).scheme}://{host}/robots.txt")
            try:
                parser.read()
            except Exception:
                parser = None  # No robots.txt reachable: treat as unrestricted.
            self._robots[host] = parser
        if parser is not None and not parser.can_fetch(USER_AGENT, url):
            return False, "disallowed by robots.txt"
        return True, ""

    def get(self, url: str, *, timeout: int = 60) -> requests.Response | None:
        ok, reason = self.allowed(url)
        if not ok:
            print(f"  skip {url} — {reason}")
            return None
        elapsed = time.time() - self._last_request
        if elapsed < REQUEST_DELAY_SECONDS:
            time.sleep(REQUEST_DELAY_SECONDS - elapsed)
        self._last_request = time.time()
        try:
            response = self.session.get(url, timeout=timeout, allow_redirects=True)
        except requests.RequestException as err:
            print(f"  skip {url} — {err}")
            return None
        if not response.ok:
            print(f"  skip {url} — HTTP {response.status_code}")
            return None
        return response


def sniff_format(url: str, content_type: str) -> str:
    suffix = Path(urlparse(url).path).suffix.lower()
    if suffix in DOCUMENT_SUFFIXES:
        return DOCUMENT_SUFFIXES[suffix]
    return CONTENT_TYPE_FORMATS.get(content_type.split(";")[0].strip().lower(), "unknown")


def write_file(directory: Path, name: str, payload: bytes) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / name
    path.write_bytes(payload)
    return path


def scrape_eurlex(fetcher: Fetcher, out_dir: Path, limit: int) -> list[Record]:
    records: list[Record] = []
    for celex in EMPLOYMENT_CELEX:
        if len(records) >= limit:
            break
        for label, extension in EURLEX_FORMATS:
            if len(records) >= limit:
                break
            url = f"https://eur-lex.europa.eu/legal-content/EN/TXT/{label}/?uri=CELEX:{celex}"
            response = fetcher.get(url)
            if response is None:
                continue
            content_type = response.headers.get("Content-Type", "")
            fmt = sniff_format(url, content_type)
            # EUR-Lex answers a missing format with an HTML error page, so a
            # response that isn't the format we asked for is a miss, not a hit.
            if label != "HTML" and fmt == "html":
                print(f"  skip {celex} {label} — served HTML instead")
                continue
            body = response.content
            if label == "HTML":
                parser = TextExtractor()
                parser.feed(body.decode("utf-8", errors="replace"))
                body = f"# CELEX {celex}\n\nSource: {url}\n\n{parser.text}".encode()
                extension, fmt = "md", "markdown"
            path = write_file(out_dir, f"{celex}.{extension}", body)
            records.append(
                Record(
                    path=str(path.relative_to(CORPUS_DIR)),
                    url=url,
                    corpus="eu-directives",
                    fmt=fmt,
                    content_type=content_type,
                    bytes=len(body),
                    sha256=hashlib.sha256(body).hexdigest(),
                    title=f"EU directive CELEX {celex}",
                    celex=celex,
                )
            )
            print(f"  ok   {celex} [{fmt}] {len(body):,} bytes")
    return records


def scrape_hr(fetcher: Fetcher, out_dir: Path, limit: int) -> list[Record]:
    records: list[Record] = []
    seen: set[str] = set()
    for page in HR_SEED_PAGES:
        if len(records) >= limit:
            break
        response = fetcher.get(page)
        if response is None:
            continue
        html = response.text
        extractor = LinkExtractor()
        extractor.feed(html)
        page_title = extractor.title or urlparse(page).path

        # The seed page itself is a usable document once flattened.
        text = TextExtractor()
        text.feed(html)
        body = f"# {page_title}\n\nSource: {page}\n\n{text.text}".encode()
        name = re.sub(r"[^a-z0-9]+", "-", urlparse(page).path.lower()).strip("-") or "index"
        path = write_file(out_dir, f"{name}.md", body)
        records.append(
            Record(
                path=str(path.relative_to(CORPUS_DIR)),
                url=page,
                corpus="hr-templates",
                fmt="markdown",
                content_type=response.headers.get("Content-Type", ""),
                bytes=len(body),
                sha256=hashlib.sha256(body).hexdigest(),
                title=page_title,
            )
        )
        print(f"  ok   {page} [markdown]")

        for href in extractor.links:
            if len(records) >= limit:
                break
            target = urljoin(page, href.split("#")[0])
            suffix = Path(urlparse(target).path).suffix.lower()
            if suffix not in DOCUMENT_SUFFIXES or target in seen:
                continue
            seen.add(target)
            doc = fetcher.get(target)
            if doc is None:
                continue
            content_type = doc.headers.get("Content-Type", "")
            fmt = sniff_format(target, content_type)
            filename = Path(urlparse(target).path).name or f"document{suffix}"
            path = write_file(out_dir, filename, doc.content)
            records.append(
                Record(
                    path=str(path.relative_to(CORPUS_DIR)),
                    url=target,
                    corpus="hr-templates",
                    fmt=fmt,
                    content_type=content_type,
                    bytes=len(doc.content),
                    sha256=hashlib.sha256(doc.content).hexdigest(),
                    title=filename,
                )
            )
            print(f"  ok   {filename} [{fmt}] {len(doc.content):,} bytes")
    return records


def synthesize_office(records: list[Record], out_dir: Path) -> list[Record]:
    """Derive .docx/.pptx from already-fetched text.

    Off by default. These files are *derived*, not scraped, and every record
    carries `derived_from` so nothing in the corpus misrepresents its origin.
    Use it only when you need format coverage the real sources didn't provide.
    """
    try:
        from docx import Document  # type: ignore[import-not-found]
        from pptx import Presentation  # type: ignore[import-not-found]
        from pptx.util import Inches  # type: ignore[import-not-found]
    except ImportError:
        print("  note: python-docx/python-pptx not installed — skipping --synthesize-office")
        return []

    derived: list[Record] = []
    sources = [r for r in records if r.fmt in {"markdown", "txt"}][:6]
    for record in sources:
        text = (CORPUS_DIR / record.path).read_text(encoding="utf-8", errors="replace")
        stem = Path(record.path).stem

        document = Document()
        document.add_heading(record.title[:200], level=1)
        for paragraph in text.split("\n\n")[:40]:
            document.add_paragraph(paragraph[:2000])
        docx_path = out_dir / f"{stem}-derived.docx"
        document.save(docx_path)

        deck = Presentation()
        for chunk in text.split("\n\n")[:10]:
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = record.title[:80]
            box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
            box.text_frame.text = chunk[:800]
        pptx_path = out_dir / f"{stem}-derived.pptx"
        deck.save(pptx_path)

        for path, fmt in ((docx_path, "docx"), (pptx_path, "pptx")):
            payload = path.read_bytes()
            derived.append(
                Record(
                    path=str(path.relative_to(CORPUS_DIR)),
                    url=record.url,
                    corpus=record.corpus,
                    fmt=fmt,
                    content_type="derived",
                    bytes=len(payload),
                    sha256=hashlib.sha256(payload).hexdigest(),
                    title=f"{record.title} (derived {fmt})",
                    derived_from=record.path,
                )
            )
        print(f"  ok   {stem} -> derived docx + pptx")
    return derived


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--limit", type=int, default=100, help="total documents across both corpora (default 100)")
    parser.add_argument("--eu-share", type=float, default=0.6, help="fraction of the budget for EU directives")
    parser.add_argument("--allow-domain", action="append", default=[], help="add a host to the allowlist (check its licence first)")
    parser.add_argument("--ignore-robots", action="store_true", help="do not consult robots.txt (don't)")
    parser.add_argument("--synthesize-office", action="store_true", help="derive docx/pptx from fetched text, clearly labelled")
    parser.add_argument("--out", type=Path, default=CORPUS_DIR)
    args = parser.parse_args()

    allowed = set(ALLOWED_HOSTS)
    for host in args.allow_domain:
        allowed.add(host)
        print(f"note: allowlisting {host} — confirm its terms permit copying before you publish this corpus")

    fetcher = Fetcher(allowed, ignore_robots=args.ignore_robots)
    eu_budget = max(1, int(args.limit * args.eu_share))
    hr_budget = max(1, args.limit - eu_budget)

    print(f"\nEU directives (budget {eu_budget})")
    records = scrape_eurlex(fetcher, args.out / "eu-directives", eu_budget)
    print(f"\nHR templates (budget {hr_budget})")
    records += scrape_hr(fetcher, args.out / "hr-templates", hr_budget)

    if args.synthesize_office:
        print("\nDerived office formats")
        records += synthesize_office(records, args.out / "hr-templates")

    manifest = args.out / "manifest.json"
    manifest.parent.mkdir(parents=True, exist_ok=True)
    manifest.write_text(json.dumps([asdict(r) for r in records], indent=2), encoding="utf-8")

    formats: dict[str, int] = {}
    for record in records:
        formats[record.fmt] = formats.get(record.fmt, 0) + 1
    print(f"\n{len(records)} documents -> {manifest}")
    print("format mix: " + ", ".join(f"{fmt}={count}" for fmt, count in sorted(formats.items())))
    if len(records) < args.limit:
        print(f"note: {args.limit - len(records)} under budget — sources returned fewer documents than requested")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
